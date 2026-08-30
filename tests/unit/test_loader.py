from pathlib import Path

import pytest
from docx import Document as DocxDocument

from src.ingest.loader import load_file


def test_load_docx_converts_headings_to_markdown(tmp_path: Path):
    document = DocxDocument()
    document.add_heading("Overview", level=1)
    document.add_paragraph("The site operates two shifts per day.")
    document.add_heading("Safety", level=2)
    document.add_paragraph("Hard hats are required in the yard.")
    path = tmp_path / "guide.docx"
    document.save(str(path))

    pages = load_file(path)

    assert len(pages) == 1
    assert pages[0].page == 1
    assert pages[0].text.splitlines() == [
        "# Overview",
        "The site operates two shifts per day.",
        "## Safety",
        "Hard hats are required in the yard.",
    ]


def test_load_docx_skips_blank_paragraphs(tmp_path: Path):
    document = DocxDocument()
    document.add_paragraph("First line.")
    document.add_paragraph("")
    document.add_paragraph("Second line.")
    path = tmp_path / "notes.docx"
    document.save(str(path))

    pages = load_file(path)

    assert pages[0].text.splitlines() == ["First line.", "Second line."]


def test_load_file_rejects_unsupported_extension(tmp_path: Path):
    path = tmp_path / "program.exe"
    path.write_bytes(b"MZ")

    with pytest.raises(ValueError):
        load_file(path)
